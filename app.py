"""
Company Website Finder — Streamlit App
=======================================
Upload an Excel file with company names, get back the same file
with a 'Website' column filled in via DuckDuckGo search.

Deploy free on Streamlit Community Cloud:
https://streamlit.io/cloud
"""

import io
import re
import time
import random

import pandas as pd
import streamlit as st
from ddgs import DDGS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ── Page config ──────────────────────────────────────────────────────

st.set_page_config(
    page_title="Company Website Finder",
    page_icon="🔍",
    layout="centered",
)

st.title("🔍 Company Website Finder")
st.markdown(
    "Upload an Excel file with company names. "
    "This tool searches the web and adds the website URL for each company."
)

# ── Skip list (directories, not real company sites) ──────────────────

SKIP = [
    "companieshouse", "gov.uk/company", "find-and-update", "endole",
    "opencorporates", "wikipedia", "linkedin.com", "facebook.com",
    "yell.com", "192.com", "crunchbase", "glassdoor", "duedil",
    "trustpilot", "amazon.", "ebay.", "indeed.com", "twitter.com",
    "youtube.com", "instagram.com", "tiktok.com", "charitycommission",
    "companies-house", "dnb.com",
]


def find_website(company_name: str) -> str:
    """Search DuckDuckGo for a company's website."""
    try:
        results = DDGS().text(
            f"{company_name} UK official website",
            region="uk-en",
            max_results=5,
        )
        for r in results:
            url = r.get("href", "")
            if url and not any(s in url.lower() for s in SKIP):
                return url
    except Exception:
        time.sleep(3)
        try:
            results = DDGS().text(
                f"{company_name} website",
                region="uk-en",
                max_results=3,
            )
            for r in results:
                url = r.get("href", "")
                if url and not any(s in url.lower() for s in SKIP):
                    return url
        except Exception:
            pass
    return ""


def find_website_short(company_name: str) -> str:
    """Retry with a shortened company name."""
    short = re.sub(
        r"\b(LIMITED|LTD|TRUST|MULTI[- ]?ACADEMY|MAT|THE|EDUCATION|PARTNERSHIP|ACADEMIES)\b",
        "", company_name, flags=re.IGNORECASE,
    ).strip()
    short = re.sub(r"\s+", " ", short).strip()
    if short and short != company_name:
        return find_website(short)
    return ""


def build_excel(df: pd.DataFrame) -> io.BytesIO:
    """Build a formatted Excel file from the dataframe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"

    cols = list(df.columns)
    hdr_font = Font(bold=True, size=10, name="Arial", color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2F5496")
    body_font = Font(size=10, name="Arial")
    link_font = Font(size=10, name="Arial", color="0563C1", underline="single")
    green_fill = PatternFill("solid", fgColor="E2EFDA")

    for c, col_name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=col_name)
        cell.font = hdr_font
        cell.fill = hdr_fill

    for r, (_, row) in enumerate(df.iterrows(), 2):
        has_site = bool(str(row.get("Website", "")).strip())
        for c, col_name in enumerate(cols, 1):
            val = str(row[col_name]) if pd.notna(row[col_name]) else ""
            if val == "nan":
                val = ""
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = body_font
            if col_name == "Website" and val:
                cell.font = link_font
            if has_site:
                cell.fill = green_fill

    for c, col_name in enumerate(cols, 1):
        w = 50 if "name" in col_name.lower() else (50 if col_name == "Website" else 18)
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ── File upload ──────────────────────────────────────────────────────

uploaded_file = st.file_uploader("Upload your Excel file (.xlsx)", type=["xlsx"])

if uploaded_file:
    df = pd.read_excel(uploaded_file, dtype=str, engine="openpyxl")
    df.columns = df.columns.str.strip()

    # Auto-detect company name column
    name_col = next(
        (c for c in df.columns if any(k in c.lower() for k in ["company", "name", "organisation"])),
        df.columns[0],
    )

    st.success(f"Loaded **{len(df)}** companies from column **\"{name_col}\"**")
    st.dataframe(df.head(5), use_container_width=True)

    # ── Run search ───────────────────────────────────────────────────

    if st.button("🔍 Find Websites", type="primary"):

        df["Website"] = ""
        total = len(df)
        found = 0

        progress_bar = st.progress(0, text="Starting...")
        status_text = st.empty()
        results_table = st.empty()

        for idx in range(total):
            company = str(df.at[idx, name_col]).strip()
            if not company or company == "nan":
                continue

            progress_bar.progress(
                (idx + 1) / total,
                text=f"Searching {idx + 1}/{total}: {company[:50]}..."
            )

            # Search
            url = find_website(company)

            # Retry with short name if needed
            if not url:
                time.sleep(random.uniform(1, 2))
                url = find_website_short(company)

            if url:
                df.at[idx, "Website"] = url
                found += 1

            status_text.markdown(
                f"**Found {found}/{idx + 1}** websites so far"
            )

            # Update preview every 10 rows
            if (idx + 1) % 10 == 0:
                results_table.dataframe(
                    df[[name_col, "Website"]].head(idx + 1),
                    use_container_width=True,
                )

            time.sleep(random.uniform(1.5, 3.0))

        progress_bar.progress(1.0, text="Done!")
        status_text.markdown(
            f"### ✅ Complete: Found **{found}/{total}** websites ({100 * found // total}%)"
        )

        # Show final results
        st.dataframe(df[[name_col, "Website"]], use_container_width=True)

        # Download button
        excel_bytes = build_excel(df)
        st.download_button(
            label="📥 Download Enriched Excel",
            data=excel_bytes,
            file_name=f"websites_{uploaded_file.name}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )

# ── Footer ───────────────────────────────────────────────────────────

st.markdown("---")
st.caption(
    "Searches via DuckDuckGo · No API keys needed · "
    "Skips directory sites (Companies House, LinkedIn, Yell, etc.)"
)
